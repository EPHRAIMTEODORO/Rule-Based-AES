"""Recovery contracts: interruption budgets, ownership, and API diagnostics."""

import argparse
import io
import json
import socket
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import Mock, patch

from HITL import hitl_processor
from HITL import local_api_server as api
from HITL import app_backend
from openpyxl import Workbook, load_workbook

scorer = hitl_processor.scorer


def response(payload=None):
    if payload is None:
        payload = {"message": {"content": '{"recommended_score": 4}'}}
    return io.BytesIO(json.dumps(payload).encode())


def interrupted():
    return urllib.error.HTTPError(
        "http://127.0.0.1:12345/api/chat", 500, "runner stopped", {},
        io.BytesIO(b'{"error":"read tcp: wsarecv: An existing connection was forcibly closed by the remote host"}'),
    )


class ChatRecoveryTests(unittest.TestCase):
    def call(self, **kwargs):
        return scorer.call_ollama_chat("rubric", "essay", "llama3:8b", "http://127.0.0.1:12345/api/chat", 0, 600, **kwargs)

    def test_runner_interruption_retries_same_essay_and_reports_recovery(self):
        recover = Mock()
        progress = Mock()
        with patch.object(scorer.urllib.request, "urlopen", side_effect=[interrupted(), response()]) as request, patch.object(scorer.time, "sleep"):
            self.assertEqual(self.call(recovery_callback=recover, progress_callback=progress), {"recommended_score": 4})
        self.assertEqual(request.call_count, 2)
        self.assertEqual(request.call_args_list[0].args[0].data, request.call_args_list[1].args[0].data)
        self.assertEqual(json.loads(request.call_args.args[0].data)["options"], {"temperature": 0})
        recover.assert_called_once()
        self.assertEqual(progress.call_args.args[0]["stage"], "recovering")

    def test_server_connection_loss_uses_recovery_callback(self):
        recover = Mock()
        with patch.object(scorer.urllib.request, "urlopen", side_effect=[urllib.error.URLError(ConnectionRefusedError()), response()]), patch.object(scorer.time, "sleep"):
            self.call(recovery_callback=recover)
        recover.assert_called_once()

    def test_recovery_stops_if_private_port_was_claimed_by_another_process(self):
        recover = Mock(side_effect=scorer.OllamaOwnershipError("private port claimed"))
        with patch.object(scorer.urllib.request, "urlopen", side_effect=interrupted()) as request, patch.object(scorer.time, "sleep"):
            with self.assertRaises(scorer.OllamaOwnershipError):
                self.call(recovery_callback=recover)
        request.assert_called_once()

    def test_shutdown_aborts_recovery_instead_of_retrying_again(self):
        recover = Mock(side_effect=scorer.OllamaRuntimeStopped("closing"))
        with patch.object(scorer.urllib.request, "urlopen", side_effect=interrupted()) as request, patch.object(scorer.time, "sleep"):
            with self.assertRaises(scorer.OllamaRuntimeStopped):
                self.call(recovery_callback=recover)
        request.assert_called_once()

    def test_connection_reset_during_response_can_recover(self):
        with patch.object(scorer.urllib.request, "urlopen", side_effect=[ConnectionResetError(), response()]), patch.object(scorer.time, "sleep"):
            self.assertEqual(self.call()["recommended_score"], 4)

    def test_repeated_interruption_stops_after_two_retries(self):
        recover = Mock()
        with patch.object(scorer.urllib.request, "urlopen", side_effect=[interrupted(), interrupted(), interrupted()]) as request, patch.object(scorer.time, "sleep"):
            with self.assertRaisesRegex(RuntimeError, "could not recover after 2 retries"):
                self.call(recovery_callback=recover)
        self.assertEqual(request.call_count, 3)
        self.assertEqual(recover.call_count, 2)

    def test_nonrecoverable_http_errors_are_not_retried(self):
        for code, message in ((404, "could not find"), (400, "rejected"), (500, "rejected")):
            with self.subTest(code=code), patch.object(scorer.urllib.request, "urlopen", side_effect=urllib.error.HTTPError("url", code, "error", {}, io.BytesIO(b'bad request'))) as request:
                with self.assertRaisesRegex(RuntimeError, message):
                    self.call()
                request.assert_called_once()

    def test_timeout_is_reported_without_restarting_or_retrying(self):
        recover = Mock()
        with patch.object(scorer.urllib.request, "urlopen", side_effect=urllib.error.URLError(socket.timeout())) as request:
            with self.assertRaisesRegex(RuntimeError, "timed out"):
                self.call(recovery_callback=recover)
        request.assert_called_once()
        recover.assert_not_called()

    def test_invalid_message_shapes_have_a_bounded_json_retry_budget(self):
        for payload in (None, [], {"message": None}, {"message": {"content": None}}, {"message": {"content": "[]"}}):
            with self.subTest(payload=payload), patch.object(scorer.urllib.request, "urlopen", side_effect=lambda *a, **k: io.BytesIO(json.dumps(payload).encode())) as request, patch.object(scorer.time, "sleep"):
                with self.assertRaisesRegex(ValueError, "valid rubric JSON"):
                    self.call()
                self.assertEqual(request.call_count, 3)

    def test_json_and_interruption_budgets_do_not_reset_each_other(self):
        attempts = [response({"message": None}), interrupted(), response({"message": None}), interrupted(), response()]
        with patch.object(scorer.urllib.request, "urlopen", side_effect=attempts) as request, patch.object(scorer.time, "sleep"):
            self.assertEqual(self.call()["recommended_score"], 4)
            self.assertEqual(request.call_count, 5)

    def test_probe_token_limit_does_not_change_normal_scoring(self):
        with patch.object(scorer.urllib.request, "urlopen", return_value=response()) as request:
            self.call(num_predict=32)
        self.assertEqual(json.loads(request.call_args.args[0].data)["options"]["num_predict"], 32)

    def test_total_request_time_budget_is_not_reset_after_an_interruption(self):
        with patch.object(scorer.time, "monotonic", side_effect=[0, 0, 599]), patch.object(scorer.urllib.request, "urlopen", side_effect=interrupted()) as request:
            with self.assertRaisesRegex(RuntimeError, "timed out.*recovery"):
                self.call()
        request.assert_called_once()


class OwnershipTests(unittest.TestCase):
    def setUp(self):
        self.stopping = patch.object(scorer, "_OLLAMA_STOPPING", threading.Event())
        self.stopping.start()
        self.addCleanup(self.stopping.stop)
        self.env = patch.dict(scorer.os.environ, {"HITL_OLLAMA_MANAGED": "0", "OLLAMA_HOST": "127.0.0.1:11434"})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.process = patch.object(scorer, "_OLLAMA_PROCESS", None)
        self.process.start()
        self.addCleanup(self.process.stop)
        self.url = patch.object(scorer, "_OLLAMA_OWNED_URL", None)
        self.url.start()
        self.addCleanup(self.url.stop)
        self.relay = patch.object(scorer, "_relay_ollama_output")
        self.relay.start()
        self.addCleanup(self.relay.stop)
        self.args = argparse.Namespace(ollama_url="http://127.0.0.1:12345/api/chat", ollama_command="bundled-ollama", ollama_startup_timeout=30, quiet=True, managed_ollama=True)

    def test_private_endpoint_never_silently_reuses_an_unowned_server(self):
        with patch.object(scorer, "is_ollama_running", return_value=True), patch.object(scorer.subprocess, "Popen") as spawn:
            with self.assertRaisesRegex(RuntimeError, "already in use"):
                scorer.start_ollama_if_needed(self.args)
        spawn.assert_not_called()

    def test_owned_server_restart_uses_the_configured_endpoint(self):
        old = Mock()
        old.poll.return_value = 1
        scorer._OLLAMA_PROCESS = old
        scorer._OLLAMA_OWNED_URL = self.args.ollama_url
        new = Mock()
        new.poll.return_value = None
        with patch.object(scorer, "is_ollama_running", side_effect=[False, True]), patch.object(scorer.subprocess, "Popen", return_value=new) as spawn:
            scorer.recover_ollama(self.args, 20)
        self.assertEqual(spawn.call_args.kwargs["env"]["OLLAMA_HOST"], "127.0.0.1:12345")
        self.assertEqual(spawn.call_args.kwargs["stdout"], scorer.subprocess.PIPE)
        self.assertEqual(spawn.call_args.kwargs["stderr"], scorer.subprocess.STDOUT)
        self.assertIs(scorer._OLLAMA_PROCESS, new)

    def test_recovery_never_spawns_or_stops_an_external_service(self):
        self.args.managed_ollama = False
        with patch.object(scorer.subprocess, "Popen") as spawn, patch.object(scorer.subprocess, "run") as kill:
            scorer.recover_ollama(self.args, 20)
            scorer.stop_ollama_if_started()
        spawn.assert_not_called()
        kill.assert_not_called()

    def test_managed_server_cannot_bind_to_a_remote_address(self):
        self.args.ollama_url = "http://example.invalid:11434/api/chat"
        with patch.object(scorer.subprocess, "Popen") as spawn:
            with self.assertRaisesRegex(RuntimeError, "loopback"):
                scorer.start_ollama_if_needed(self.args)
        spawn.assert_not_called()

    def test_shutdown_prevents_a_concurrent_retry_from_restarting_the_server(self):
        scorer.stop_ollama_if_started()
        with patch.object(scorer.subprocess, "Popen") as spawn:
            with self.assertRaises(scorer.OllamaRuntimeStopped):
                scorer.recover_ollama(self.args, 20)
        spawn.assert_not_called()

    def test_shutdown_targets_only_the_owned_pid(self):
        process = Mock(pid=98765)
        process.poll.return_value = None
        scorer._OLLAMA_PROCESS = process
        with patch.object(scorer.subprocess, "run") as stop:
            scorer.stop_ollama_if_started()
        if scorer.os.name == "nt":
            self.assertEqual(stop.call_args.args[0], ["taskkill", "/PID", "98765", "/T", "/F"])
        else:
            process.terminate.assert_called_once()
        self.assertIsNone(scorer._OLLAMA_PROCESS)


class RecoveryWorkflowTests(unittest.TestCase):
    def test_hidden_child_output_is_relayed_into_the_backend_log(self):
        process = Mock(pid=123, stdout=io.StringIO("runner error details\n"))
        with self.assertLogs(scorer.LOGGER, level="INFO") as logs:
            scorer._relay_ollama_output(process)
        self.assertIn("runner error details", "\n".join(logs.output))
        self.assertTrue(process.stdout.closed)

    def test_recovery_on_later_essay_keeps_completed_rows_and_progress(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "input.xlsx"
            output = Path(folder) / "output.xlsx"
            workbook = Workbook()
            workbook.active.append(["Student ID", "Topic", "Essay"])
            for number in (1, 2):
                workbook.active.append([str(number), "topic", "essay"])
            workbook.save(source)
            events = []
            calls = []

            def evaluate(row, args, index, text_column, progress_callback=None):
                calls.append(index)
                if index == 2:
                    progress_callback({"stage": "recovering", "message": "Recovering this essay"})
                return {**row, "llm_recommended_score": 4, "aes_grammar_errors_per_100": None}

            with patch.object(scorer, "evaluate_row", side_effect=evaluate):
                result = hitl_processor.process_workbook(str(source), str(output), start_ollama=False, progress_callback=events.append)
            self.assertEqual(calls, [1, 2])
            self.assertEqual(result.rows_processed, 2)
            recovery = next(event for event in events if event["stage"] == "recovering")
            self.assertEqual((recovery["current"], recovery["rows_processed"]), (2, 1))
            self.assertTrue(all(record["Decision_Status"] == "Pending" for record in result.records))
            saved = load_workbook(output, read_only=True)
            self.assertEqual(saved.active.max_row, 3)
            saved.close()

    def test_failure_records_and_logs_original_traceback(self):
        job = app_backend.AppJob(job_id="trace-test", input_path="input.xlsx", output_path="output.xlsx")
        with patch.dict(app_backend._JOBS, {job.job_id: job}), patch.object(app_backend, "process_workbook", side_effect=RuntimeError("runner interrupted")):
            with self.assertLogs("HITL.app_backend", level="ERROR") as captured:
                app_backend._run_job(job.job_id, {})
            self.assertIn("runner interrupted", job.traceback)
            self.assertEqual(job.status, "failed")
            self.assertIn("Traceback", "\n".join(captured.output))


class InferenceApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), api.HITLRequestHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f"http://127.0.0.1:{cls.server.server_port}/inference-check"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def test_deliberate_check_runs_inference_and_returns_runtime_provenance(self):
        runtime = {"app_version": "0.1.3", "mode": "app-owned", "ollama_version": "test"}
        with patch.object(api, "list_jobs", return_value=[]), patch.object(api.scorer, "call_ollama_chat", return_value={"ok": True}) as call, patch.object(api, "runtime_info", return_value=runtime):
            with urllib.request.urlopen(urllib.request.Request(self.url, data=b"")) as result:
                payload = json.load(result)
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["runtime"], runtime)
        self.assertEqual(call.call_args.kwargs["num_predict"], 32)

    def test_inference_check_does_not_run_during_workbook_processing(self):
        with patch.object(api, "list_jobs", return_value=[{"status": "running"}]), patch.object(api.scorer, "call_ollama_chat") as call:
            with self.assertRaises(urllib.error.HTTPError) as error:
                urllib.request.urlopen(urllib.request.Request(self.url, data=b""))
            self.assertEqual(error.exception.code, 409)
        call.assert_not_called()


if __name__ == "__main__":
    unittest.main()
